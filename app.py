import os
import json
import re
import tempfile
from typing import Dict, Any, List

import streamlit as st

import google.generativeai as genai

import pdfplumber
import docx


SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
}


def extract_text(file_bytes: bytes, file_type: str) -> str:
    """Extract resume text from uploaded content."""
    if file_type == "pdf":
        with pdfplumber.open(io_bytes=file_bytes) as pdf:
            texts = []
            for page in pdf.pages:
                t = page.extract_text() or ""
                texts.append(t)
            return "\n".join(texts).strip()

    if file_type == "docx":
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        doc = docx.Document(tmp_path)
        return "\n".join([p.text for p in doc.paragraphs]).strip()

    if file_type == "txt":
        return file_bytes.decode("utf-8", errors="ignore").strip()

    raise ValueError(f"Unsupported file_type: {file_type}")


def safe_truncate(text: str, max_chars: int = 12000) -> str:
    """Limit size for model context safety."""
    text = text.strip()
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[TRUNCATED]"


def build_prompt(resume_text: str, job_title: str = "", target_keywords: str = "") -> str:
    extra = []
    if job_title.strip():
        extra.append(f"Job Title: {job_title.strip()}")
    if target_keywords.strip():
        extra.append(f"Target Keywords/Role Requirements: {target_keywords.strip()}")

    extra_block = "\n".join(extra).strip()
    extra_block = f"\nContext:\n{extra_block}\n" if extra_block else ""

    # Instruct Gemini to output strict JSON.
    return f"""
You are an ATS (Applicant Tracking System) resume evaluator.
Evaluate the resume using ATS best practices and output ONLY valid JSON.

{extra_block}
Resume Text:
\"\"\"
{resume_text}
\"\"\"

Return JSON with this exact schema:
{{
  "ats_score": number, 
  "score_breakdown": {{
    "formatting": number,
    "keywords": number,
    "experience_relevance": number,
    "clarity_structure": number,
    "achievements": number
  }},
  "summary": string,
  "improvements": [
    {{
      "category": string,
      "issue": string,
      "recommendation": string,
      "example_fix": string
    }}
  ],
  "rewrite_suggestions": [
    {{
      "section": string,
      "suggested_text": string
    }}
  ]
}}

Rules:
- ats_score must be an integer from 0 to 100.
- improvements array must have at least 6 items.
- rewrite_suggestions can be empty but if provided, include at least 2 items.
- Ensure text recommendations are specific to the provided resume.
- Output ONLY JSON. No markdown. No extra keys.
""".strip()


def generate_ats_report(resume_text: str, job_title: str, target_keywords: str) -> Dict[str, Any]:
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise RuntimeError("Missing GOOGLE_API_KEY environment variable.")

    genai.configure(api_key=api_key)

    # Gemini Flash model
    model = genai.GenerativeModel("gemini-1.5-flash")

    prompt = build_prompt(
        resume_text=safe_truncate(resume_text),
        job_title=job_title,
        target_keywords=target_keywords,
    )

    # Ask for JSON response
    response = model.generate_content(
        prompt,
        generation_config=genai.types.GenerationConfig(
            temperature=0.2,
            max_output_tokens=1200,
        ),
    )

    raw = response.text.strip()

    # Try direct JSON parse first
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        # Fallback: extract JSON object from text
        m = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if not m:
            raise ValueError("Model did not return valid JSON.")
        return json.loads(m.group(0))


# --- Streamlit UI ---
st.set_page_config(page_title="ATS Resume Scorer", layout="wide")
st.title("📄 ATS Resume Scorer (Gemini Flash)")

with st.sidebar:
    st.header("⚙️ Settings")
    job_title = st.text_input("Target Job Title (optional)", value="")
    target_keywords = st.text_area(
        "Target Keywords / Role Requirements (optional)",
        value="",
        height=120
    )

    st.markdown("---")
    st.info("Set environment variable `GOOGLE_API_KEY` to enable Gemini.")

uploaded = st.file_uploader(
    "Upload your resume (PDF, DOCX, or TXT)",
    type=["pdf", "docx", "txt"]
)

if uploaded is not None:
    file_type = SUPPORTED_TYPES.get(uploaded.type)
    if not file_type:
        st.error(f"Unsupported file type: {uploaded.type}")
        st.stop()

    file_bytes = uploaded.read()
    if not file_bytes:
        st.error("Uploaded file is empty.")
        st.stop()

    # Extract text
    with st.spinner("Extracting resume text..."):
        try:
            # pdfplumber expects a file path or file-like; pdfplumber.open supports file-like
            # We'll create a BytesIO for consistency.
            import io
            if file_type == "pdf":
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    texts = []
                    for page in pdf.pages:
                        t = page.extract_text() or ""
                        texts.append(t)
                    resume_text = "\n".join(texts).strip()
            elif file_type == "docx":
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                    tmp.write(file_bytes)
                    tmp_path = tmp.name
                doc = docx.Document(tmp_path)
                resume_text = "\n".join([p.text for p in doc.paragraphs]).strip()
            else:
                resume_text = file_bytes.decode("utf-8", errors="ignore").strip()
        except Exception as e:
            st.error(f"Text extraction failed: {e}")
            st.stop()

    if not resume_text or len(resume_text) < 50:
        st.error("Could not extract enough text from the resume. Please upload a clearer PDF/DOCX.")
        st.stop()

    with st.expander("🔎 Preview extracted text"):
        st.text_area("Extracted Resume Text", resume_text[:5000] + ("..." if len(resume_text) > 5000 else ""), height=260)

    if st.button("Analyze Resume", type="primary"):
        with st.spinner("Analyzing with Gemini Flash..."):
            try:
                result = generate_ats_report(
                    resume_text=resume_text,
                    job_title=job_title,
                    target_keywords=target_keywords,
                )
            except Exception as e:
                st.error(f"AI analysis failed: {e}")
                st.stop()

        # --- Render results ---
        ats_score = result.get("ats_score", 0)
        st.subheader("📈 ATS Score")

        col1, col2, col3 = st.columns([2, 1, 2])
        with col1:
            st.metric("ATS Score", int(ats_score))
        with col2:
            st.write("")

        with col3:
            breakdown = result.get("score_breakdown", {})
            st.write("**Score Breakdown**")
            for k, v in breakdown.items():
                st.write(f"- {k}: {v}")

        st.subheader("🧠 Summary")
        st.write(result.get("summary", ""))

        st.subheader("✅ Improvements")
        improvements: List[Dict[str, str]] = result.get("improvements", [])
        for i, imp in enumerate(improvements, start=1):
            category = imp.get("category", "General")
            issue = imp.get("issue", "")
            rec = imp.get("recommendation", "")
            example = imp.get("example_fix", "")

            with st.expander(f"{i}. {category}"):
                st.write("**Issue:** " + issue)
                st.write("**Recommendation:** " + rec)
                if example:
                    st.write("**Example Fix:**")
                    st.code(example, language="markdown")

        st.subheader("✍️ Suggested Rewrite (if available)")
        rewrites = result.get("rewrite_suggestions", [])
        if not rewrites:
            st.info("No rewrite suggestions provided by the model.")
        else:
            for r in rewrites:
                section = r.get("section", "Section")
                suggested_text = r.get("suggested_text", "")
                with st.expander(section):
                    st.code(suggested_text, language="markdown")
